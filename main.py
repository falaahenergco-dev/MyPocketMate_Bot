import os
import io
import json
import time
import logging
import tempfile
import threading
from datetime import datetime
from http.server import HTTPServer, BaseHTTPRequestHandler

import openpyxl
from openpyxl.comments import Comment
from PIL import Image

from telegram import Update
from telegram.ext import ApplicationBuilder, CommandHandler, MessageHandler, filters, ContextTypes

from google import genai
from google.genai import types
from google.oauth2.service_account import Credentials
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload, MediaFileUpload

logging.basicConfig(format='%(asctime)s - %(name)s - %(levelname)s - %(message)s', level=logging.INFO)
logging.getLogger("httpx").setLevel(logging.WARNING)  # حتى لا يظهر توكن البوت في السجلات
logger = logging.getLogger(__name__)

TELEGRAM_TOKEN = os.environ.get("TELEGRAM_TOKEN")
GEMINI_API_KEY = os.environ.get("GEMINI_API_KEY")
SERVICE_ACCOUNT_JSON = os.environ.get("SERVICE_ACCOUNT_JSON")
GOOGLE_FILE_ID = os.environ.get("GOOGLE_FILE_ID", "1jf8dawjd96Cxi2H7Ik6st4R2rzsDH8Gd")
GEMINI_MODEL = os.environ.get("GEMINI_MODEL", "gemini-3.8-flash")

SHEET_NAME = "المصروفات والمدخرات"
FIRST_DATA_ROW = 2  # الصف 1 فيه العناوين

# أعمدة الملف (حسب ملف المصاريف والمدخرات.xlsx)
COLUMNS = {
    'الكافيهات': 2,          # B
    'الجامعة': 3,            # C
    'الجيم': 4,              # D
    'مقتنيات ذهبية': 5,      # E
    'القروض (Loans)': 6,     # F
    'هدايا/دخل إضافي': 7,    # G
    'المدخرات الثابتة': 8,   # H
}


# ---------- خادم صغير لـ Render (فحص المنفذ) ----------
class Health(BaseHTTPRequestHandler):
    def do_GET(self):
        self.send_response(200)
        self.end_headers()
        self.wfile.write(b"Bot is running!")

    def do_HEAD(self):
        self.send_response(200)
        self.end_headers()

    def log_message(self, *a):
        pass


def run_health_server():
    port = int(os.environ.get("PORT", 10000))
    HTTPServer(("0.0.0.0", port), Health).serve_forever()


# ---------- Gemini ----------
SYSTEM_PROMPT = """You are a personal finance assistant. Analyze Arabic text or a receipt image.

Categories (use the EXACT Arabic name):
- 'الكافيهات': coffee, cafes, restaurants, drinks, snacks, food
- 'الجامعة': university, books, tuition, study, transport to university
- 'الجيم': gym, fitness, sports, supplements
- 'مقتنيات ذهبية': gold, jewelry, assets
- 'القروض (Loans)': loans, debts, borrowed or lent money
- 'هدايا/دخل إضافي': gifts, bonus, side income, money received
- 'المدخرات الثابتة': fixed savings, deposits, money saved

Actions:
1. Record an amount:
   {"action":"add","data":{"category":"exact_category_name","amount":NUMBER,"details":"short description"}}
2. Question about totals (كم صرفت / مجموع / شكد):
   {"action":"query","data":{"category":"exact_category_name or ALL","period":"month or all"}}
3. Not related to money or unclear:
   {"action":"unknown","data":{}}

Amounts: convert words to numbers (مليون = 1000000, ألف = 1000, ربع مليون = 250000).
If a receipt has several items of the same category, return ONE item with the total.
Return ONLY a valid JSON array. No markdown."""


def call_gemini(contents):
    """يرسل الطلب لـ Gemini ويعيد المحاولة تلقائياً عند الازدحام المؤقت (503/429)."""
    client = genai.Client(api_key=GEMINI_API_KEY)
    last_err = None
    for attempt in range(4):
        try:
            resp = client.models.generate_content(
                model=GEMINI_MODEL,
                contents=contents,
                config=types.GenerateContentConfig(response_mime_type="application/json"),
            )
            raw = resp.text.strip().replace("```json", "").replace("```", "").strip()
            data = json.loads(raw)
            return data if isinstance(data, list) else [data]
        except Exception as e:
            last_err = e
            if any(c in str(e) for c in ("503", "429", "UNAVAILABLE", "RESOURCE_EXHAUSTED")):
                time.sleep(2 * (attempt + 1))
                continue
            raise
    raise last_err


# ---------- Google Drive (نفس طريقة بوت الكوزمتك) ----------
def get_drive_service():
    info = json.loads(SERVICE_ACCOUNT_JSON)
    creds = Credentials.from_service_account_info(info, scopes=["https://www.googleapis.com/auth/drive"])
    return build("drive", "v3", credentials=creds)


def download_excel():
    service = get_drive_service()
    req = service.files().get_media(fileId=GOOGLE_FILE_ID)
    buf = io.BytesIO()
    dl = MediaIoBaseDownload(buf, req)
    done = False
    while not done:
        _, done = dl.next_chunk()
    buf.seek(0)
    return openpyxl.load_workbook(buf)


def upload_excel(wb):
    out = io.BytesIO()
    wb.save(out)
    tmp_path = os.path.join(tempfile.gettempdir(), "expenses_upload.xlsx")
    with open(tmp_path, "wb") as f:
        f.write(out.getvalue())
    media = MediaFileUpload(
        tmp_path, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    get_drive_service().files().update(fileId=GOOGLE_FILE_ID, media_body=media).execute()
    try:
        os.unlink(tmp_path)
    except OSError:
        pass


# ---------- Excel ----------
def get_sheet(wb):
    return wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.worksheets[0]


def find_empty_row(ws):
    r = FIRST_DATA_ROW
    while any(ws.cell(row=r, column=c).value not in (None, "") for c in range(1, 9)):
        r += 1
    return r


def to_number(x):
    try:
        return float(str(x).replace(",", ""))
    except (TypeError, ValueError):
        return 0.0


def fmt(n):
    n = float(n)
    return f"{int(n):,}" if n.is_integer() else f"{n:,.2f}"


def add_entry(ws, category, amount, details=""):
    col = COLUMNS.get(category)
    if col is None:
        return f"⚠️ لم أعرف التصنيف: {category}"
    amount = to_number(amount)
    if amount <= 0:
        return "⚠️ لم أجد مبلغاً صحيحاً في الرسالة"

    row = find_empty_row(ws)
    ws.cell(row=row, column=1, value=datetime.now().strftime("%Y-%m-%d"))
    cell = ws.cell(row=row, column=col, value=amount)
    if details:
        cell.comment = Comment(details, "PocketMate")
    return f"✅ تم تسجيل {fmt(amount)} في [{category}]" + (f"\n📝 {details}" if details else "")


def query_totals(ws, category="ALL", period="month"):
    this_month = datetime.now().strftime("%Y-%m")
    totals = {name: 0.0 for name in COLUMNS}
    r = FIRST_DATA_ROW
    while r <= ws.max_row:
        date_val = ws.cell(row=r, column=1).value
        if date_val:
            date_str = date_val.strftime("%Y-%m-%d") if isinstance(date_val, datetime) else str(date_val)
            if period == "all" or date_str.startswith(this_month):
                for name, col in COLUMNS.items():
                    totals[name] += to_number(ws.cell(row=r, column=col).value)
        r += 1

    label = "هذا الشهر" if period != "all" else "الكلي"
    if category in COLUMNS:
        return f"📊 مجموع [{category}] {label}: {fmt(totals[category])}"
    lines = [f"📊 المجاميع {label}:"]
    lines += [f"• {name}: {fmt(v)}" for name, v in totals.items() if v]
    return "\n".join(lines) if len(lines) > 1 else f"📊 لا توجد مبالغ مسجلة {label}"


def process(items):
    """ينفّذ أوامر Gemini على ملف الإكسل ويعيد نص الرد."""
    adds = [i for i in items if i.get("action") == "add"]
    queries = [i for i in items if i.get("action") == "query"]
    if not adds and not queries:
        return "🤔 لم أفهم. اكتب مثلاً: قهوة 5000 أو كم صرفت على الكافيهات هذا الشهر؟"

    wb = download_excel()
    ws = get_sheet(wb)
    replies = []
    for i in adds:
        d = i.get("data", {})
        replies.append(add_entry(ws, d.get("category"), d.get("amount"), d.get("details", "")))
    if adds:
        upload_excel(wb)
    for q in queries:
        d = q.get("data", {})
        replies.append(query_totals(ws, d.get("category", "ALL"), d.get("period", "month")))
    return "\n\n".join(replies)


# ---------- Telegram ----------
async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text(
        "👋 أهلاً بك! أنا بوت تتبع المصاريف والمدخرات.\n\n"
        "📝 اكتب مثلاً:\n"
        "• قهوة 5000\n"
        "• دفعت اشتراك الجيم 40 ألف\n"
        "• ادخرت 100 ألف\n"
        "• كم صرفت على الكافيهات هذا الشهر؟\n\n"
        "📸 أو أرسل صورة فاتورة"
    )


async def handle_message(update: Update, context: ContextTypes.DEFAULT_TYPE):
    msg = await update.message.reply_text("⏳ جاري المعالجة...")
    try:
        items = call_gemini([SYSTEM_PROMPT, f"User: {update.message.text}"])
        await msg.edit_text(process(items))
    except Exception as e:
        logger.exception("Text error")
        await msg.edit_text(f"❌ خطأ: {type(e).__name__}: {str(e)[:300]}")


async def handle_photo(update: Update, context: ContextTypes.DEFAULT_TYPE):
    msg = await update.message.reply_text("📸 جاري تحليل الفاتورة...")
    try:
        photo_file = await update.message.photo[-1].get_file()
        image = Image.open(io.BytesIO(bytes(await photo_file.download_as_bytearray())))
        items = call_gemini([SYSTEM_PROMPT, image, "Extract the expenses from this receipt."])
        await msg.edit_text(process(items))
    except Exception as e:
        logger.exception("Photo error")
        await msg.edit_text(f"❌ خطأ: {type(e).__name__}: {str(e)[:300]}")


def main():
    threading.Thread(target=run_health_server, daemon=True).start()

    app = ApplicationBuilder().token(TELEGRAM_TOKEN).build()
    app.add_handler(CommandHandler("start", start))
    app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, handle_message))
    app.add_handler(MessageHandler(filters.PHOTO, handle_photo))
    logger.info("🤖 PocketMate يعمل!")
    app.run_polling()


if __name__ == '__main__':
    main()
